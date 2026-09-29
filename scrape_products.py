"""
chemifriend.com의 Product 게시판(bo_table=product) 전체를 읽어서
content/content.xlsx 의 "제품상세" 시트에 자동으로 채워 넣는 스크립트입니다.

왜 이 스크립트가 필요한가:
- 기존 사이트의 Product 게시판에는 카테고리(CABOT/SYNTHOMER-Eastman/ARKEMA/SOLVAY/Others) 아래
  18개 게시글이 있고, 그 안에 다시 소그룹(예: "EASTOFLEX > Propylene Homopolymers")별로
  스펙 표가 중첩되어 있습니다. 사람이 손으로 옮기면 실수가 나기 쉬워서, 실제 사이트에서
  직접 읽어와 자동으로 옮깁니다.

사용법 (최초 1회, 혹은 원본 사이트가 살아있는 동안 다시 실행 가능):
    pip install requests beautifulsoup4 lxml
    python scrape_products.py

실행하면:
    1) chemifriend.com에서 Product 카테고리 목록을 찾고
    2) 카테고리별 게시글 목록(wr_id)을 찾고
    3) 게시글마다 들어있는 표를 전부 읽어서
    4) content/content.xlsx 의 "제품상세" 시트를 새로 만들어 채웁니다.
       (기존 "회사정보/연혁/조직도/사업영역" 시트는 그대로 둡니다)

실행 후 반드시 content.xlsx를 열어서 내용이 잘 들어왔는지 한 번 확인해주세요.
표 구조가 특이한 페이지(예: 소그룹이 3중으로 중첩된 경우)는 완벽하지 않을 수 있습니다.
"""
import re
import time
from pathlib import Path

import requests
from bs4 import BeautifulSoup, Tag
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

BASE = "http://www.chemifriend.com"
LIST_URL = f"{BASE}/bbs/board.php?bo_table=product"
CONTENT_XLSX = Path(__file__).parent / "content" / "content.xlsx"
HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}


def get_soup(url: str) -> BeautifulSoup:
    resp = requests.get(url, headers=HEADERS, timeout=15)
    resp.encoding = resp.apparent_encoding or "utf-8"
    return BeautifulSoup(resp.text, "lxml")


def find_categories() -> list[str]:
    """Product 목록 페이지 좌측/상단 카테고리 메뉴에서 sca 값들을 찾는다."""
    soup = get_soup(LIST_URL)
    cats = []
    for a in soup.find_all("a", href=True):
        m = re.search(r"[?&]sca=([^&]+)", a["href"])
        if m:
            val = m.group(1)
            if val not in cats:
                cats.append(val)
    return cats or ["CABOT", "SYNTHOMER-Eastman", "ARKEMA", "SOLVAY", "Others"]


def find_posts_in_category(sca: str) -> list[tuple[str, str]]:
    """카테고리 안의 (wr_id, 제목) 목록을 찾는다."""
    url = f"{LIST_URL}&sca={sca}"
    soup = get_soup(url)
    posts = []
    seen = set()
    for a in soup.find_all("a", href=True):
        m = re.search(r"[?&]wr_id=(\d+)", a["href"])
        if not m:
            continue
        wr_id = m.group(1)
        title = a.get_text(strip=True)
        # 사이트 링크 텍스트가 "CABOTCarbon Black"처럼 카테고리명+제목이 붙어서 나오는 경우 분리
        if title.upper().startswith(sca.upper()):
            title = title[len(sca):].strip() or title
        if not title or wr_id in seen:
            continue
        seen.add(wr_id)
        posts.append((wr_id, title))
    return posts


def is_junk_table(table: Tag) -> bool:
    """상단 네비게이션 등 콘텐츠와 무관한 표를 걸러낸다."""
    text = table.get_text(" ", strip=True)
    if not text:
        return True
    junk_markers = ["Product", "Contact"]
    if len(text) < 30 and all(m in text for m in junk_markers):
        return True
    if "회원아이디" in text or "검색대상" in text:
        return True
    return False


def _clean_ws(s: str) -> str:
    """줄바꿈/중복 공백을 정리한다."""
    return re.sub(r"\s+", " ", s).strip() if s else s


def parse_detail_page(wr_id: str, sca: str) -> list[dict]:
    """
    게시글 하나(스펙 표 여러 개 포함 가능)를 읽어서
    [{"소그룹": ..., "품명": ..., "용도": ..., "제품형태": ..., "기타스펙": ...}, ...] 형태로 반환.
    """
    url = f"{LIST_URL}&wr_id={wr_id}&sca={sca}"
    soup = get_soup(url)

    h1 = soup.find("h1")
    list_link = soup.find("a", string=re.compile("목록"))

    # 본문 범위: <h1> 이후 ~ "목록" 링크 이전까지
    body_nodes = []
    if h1:
        for el in h1.find_all_next():
            if list_link is not None and el is list_link:
                break
            body_nodes.append(el)
    else:
        body_nodes = soup.find_all(True)

    rows_out = []
    current_label_parts: list[str] = []

    for el in body_nodes:
        if not isinstance(el, Tag):
            continue

        if el.name in ("strong", "b", "h2", "h3", "h4"):
            # 표 안에 있는 strong/b는 라벨이 아니라 데이터이므로 제외
            if el.find_parent("table") is not None:
                continue
            label_text = el.get_text(" ", strip=True)
            if label_text and len(label_text) < 60:
                current_label_parts.append(label_text)

        elif el.name == "table":
            if is_junk_table(el):
                continue
            trs = el.find_all("tr")
            if len(trs) < 2:
                continue
            header_cells = [c.get_text(" ", strip=True) for c in trs[0].find_all(["th", "td"])]
            if not header_cells:
                continue

            label = " / ".join(current_label_parts[-2:]) if current_label_parts else ""
            label = re.sub(r"^(Product|Pruduct)\s*/\s*", "", label, flags=re.I)
            current_label_parts = []  # 표 하나 처리했으면 라벨 소비

            # 용도/제품형태에 해당하는 컬럼 인덱스 찾기 (표기 흔들림 대응)
            def find_col(*keywords):
                for i, h in enumerate(header_cells):
                    for kw in keywords:
                        if kw.lower() in h.lower():
                            return i
                return None

            name_idx = 0
            app_idx = find_col("application", "용도")
            form_idx = find_col("form", "형태")

            for tr in trs[1:]:
                cells = [c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])]
                if not cells or not any(cells):
                    continue
                cells += [""] * (len(header_cells) - len(cells))

                name = cells[name_idx] if name_idx < len(cells) else ""
                application = cells[app_idx] if app_idx is not None and app_idx < len(cells) else ""
                form = cells[form_idx] if form_idx is not None and form_idx < len(cells) else ""
                name, application, form = _clean_ws(name), _clean_ws(application), _clean_ws(form)

                etc_parts = []
                for i, h in enumerate(header_cells):
                    if i in (name_idx, app_idx, form_idx):
                        continue
                    if i < len(cells) and cells[i]:
                        h_clean = _clean_ws(h).replace("I 2 No.", "I2No.").replace("I 2No.", "I2No.")
                        etc_parts.append(f"{h_clean}: {_clean_ws(cells[i])}")

                rows_out.append({
                    "소그룹": label,
                    "품명": name,
                    "용도": application,
                    "제품형태": form,
                    "기타스펙": "; ".join(etc_parts),
                })

    return rows_out


def write_to_excel(all_rows: list[dict]):
    if not CONTENT_XLSX.exists():
        raise SystemExit(f"content.xlsx를 찾을 수 없습니다: {CONTENT_XLSX}")

    wb = openpyxl.load_workbook(CONTENT_XLSX)
    if "제품상세" in wb.sheetnames:
        del wb["제품상세"]
    ws = wb.create_sheet("제품상세")

    headers = ["카테고리", "페이지명", "소그룹", "품명", "용도", "제품형태", "기타스펙"]
    ws.append(headers)
    for r in all_rows:
        ws.append([r.get(h, "") for h in headers])

    widths = [16, 26, 30, 20, 30, 14, 60]
    for col, w in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = w
    for c in [f"{col}1" for col in "ABCDEFG"]:
        ws[c].font = Font(bold=True, color="FFFFFF")
        ws[c].fill = PatternFill("solid", fgColor="1F3864")
    for row in ws.iter_rows(min_row=2):
        row[6].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(CONTENT_XLSX)


def main():
    print("1) 카테고리 찾는 중...")
    categories = find_categories()
    print("   →", categories)

    all_rows = []
    for sca in categories:
        print(f"2) '{sca}' 카테고리 게시글 목록 확인 중...")
        posts = find_posts_in_category(sca)
        for wr_id, title in posts:
            print(f"   - [{sca}] {title} (wr_id={wr_id}) 읽는 중...")
            try:
                rows = parse_detail_page(wr_id, sca)
            except Exception as e:
                print(f"     ! 실패: {e}")
                continue
            for r in rows:
                r["카테고리"] = sca
                r["페이지명"] = title
            all_rows.extend(rows)
            time.sleep(0.3)  # 서버 부담을 주지 않기 위한 최소 대기

    print(f"3) 총 {len(all_rows)}개 품목 수집 완료. content.xlsx에 저장 중...")
    write_to_excel(all_rows)
    print("완료! content/content.xlsx의 '제품상세' 시트를 확인해보세요.")


if __name__ == "__main__":
    main()
