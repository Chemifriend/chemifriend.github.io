"""
content/content.xlsx 파일의 내용을 읽어서 docs/ 폴더에 정적 HTML을 생성합니다.

사용법:
    python build.py

수정 방법:
    1. content/content.xlsx 파일을 엑셀로 열어서 내용만 고칩니다. (구조/시트명은 바꾸지 마세요)
    2. 저장합니다.
    3. deploy.bat 더블클릭 → GitHub에 push → GitHub Actions가 이 파일을 실행해 사이트를 배포합니다.
       (로컬에서 결과만 미리 보려면 이 파일을 실행하고 docs/index.html 을 브라우저로 열어보세요)

사이트 구조:
    docs/index.html                          메인 페이지
    docs/product.html                        Product 허브 (브랜드 목록)
    docs/product/<브랜드>/index.html          브랜드 소개 + 제품군 목록
    docs/product/<브랜드>/<제품군>/index.html  제품군 스펙표
    docs/product/<브랜드>/<제품군>/<품명>.html 개별 제품(SKU) 페이지
    docs/contact.html                        Contact 페이지
"""
import os
import re
import shutil
import stat
import time
from pathlib import Path
from collections import OrderedDict

import openpyxl
from jinja2 import Environment, FileSystemLoader

ROOT = Path(__file__).parent
CONTENT_XLSX = ROOT / "content" / "content.xlsx"
TEMPLATES_DIR = ROOT / "templates"
DOCS_DIR = ROOT / "docs"


def _force_remove_readonly(func, path, exc_info):
    """윈도우/OneDrive에서 읽기전용·동기화 잠금으로 삭제가 막힐 때 재시도한다."""
    try:
        os.chmod(path, stat.S_IWRITE)
        func(path)
    except Exception:
        pass


def safe_rmtree(path: Path, retries: int = 5, delay: float = 0.5):
    for attempt in range(retries):
        if not path.exists():
            return
        shutil.rmtree(path, onerror=_force_remove_readonly)
        if not path.exists():
            return
        time.sleep(delay)


def slugify(text: str) -> str:
    text = re.sub(r"[^a-zA-Z0-9가-힣]+", "-", text).strip("-")
    return text.lower() or "item"


def read_keyvalue_sheet(ws):
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        key, value = row[0], row[1]
        data[str(key).strip()] = "" if value is None else str(value)
    return data


def read_table_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        if not any(row):
            continue
        item = {headers[i]: ("" if v is None else str(v)) for i, v in enumerate(row) if i < len(headers)}
        result.append(item)
    return result


def parse_spec(spec_str: str) -> "OrderedDict[str, str]":
    """'키: 값; 키: 값' 형태의 문자열을 순서를 보존한 dict로 변환한다."""
    result = OrderedDict()
    if not spec_str:
        return result
    for part in spec_str.split(";"):
        part = part.strip()
        if not part:
            continue
        if ":" in part:
            k, v = part.split(":", 1)
            result[k.strip()] = v.strip()
        else:
            result[part] = ""
    return result


def load_content():
    wb = openpyxl.load_workbook(CONTENT_XLSX, data_only=True)
    company = read_keyvalue_sheet(wb["회사정보"])
    history = read_table_sheet(wb["연혁"])
    business_areas = read_table_sheet(wb["사업영역"])
    area_by_slug = {a.get("슬러그", ""): a for a in business_areas}

    brand_intro_rows = read_table_sheet(wb["브랜드소개"]) if "브랜드소개" in wb.sheetnames else []
    brand_intro = {b.get("슬러그", ""): b for b in brand_intro_rows}

    dept_rows = read_table_sheet(wb["부서"])
    people_rows = read_table_sheet(wb["조직원"])
    org_by_dept = OrderedDict()
    for d in dept_rows:
        org_by_dept[d.get("부서", "")] = {"설명": d.get("설명", ""), "인원": []}
    for p in people_rows:
        dept = p.get("부서", "")
        org_by_dept.setdefault(dept, {"설명": "", "인원": []})
        org_by_dept[dept]["인원"].append(p)

    # 제품상세 -> 카테고리 > 페이지명(제품군) > 소그룹 > 품목 리스트, 그룹별 동적 스펙 컬럼 계산
    products_by_category = OrderedDict()
    if "제품상세" in wb.sheetnames:
        detail_rows = read_table_sheet(wb["제품상세"])
        raw = OrderedDict()
        for row in detail_rows:
            cat = row.get("카테고리", "")
            page = row.get("페이지명", "")
            label = row.get("소그룹", "")
            raw.setdefault(cat, OrderedDict())
            raw[cat].setdefault(page, OrderedDict())
            raw[cat][page].setdefault(label, []).append(row)

        for cat, pages in raw.items():
            products_by_category[cat] = OrderedDict()
            for page, groups in pages.items():
                rendered_groups = []
                for label, items in groups.items():
                    spec_columns = []  # 순서를 보존한 컬럼 목록
                    rendered_items = []
                    seen_slugs = {}
                    for item in items:
                        spec = parse_spec(item.get("기타스펙", ""))
                        form = item.get("제품형태", "")
                        if form:
                            spec["형태"] = form
                        for k in spec:
                            if k not in spec_columns:
                                spec_columns.append(k)
                        name = item.get("품명", "")
                        base_slug = slugify(name)
                        n = seen_slugs.get(base_slug, 0) + 1
                        seen_slugs[base_slug] = n
                        sku_slug = base_slug if n == 1 else f"{base_slug}-{n}"
                        rendered_items.append({
                            "품명": name,
                            "용도": item.get("용도", ""),
                            "spec": spec,
                            "sku_slug": sku_slug,
                        })
                    rendered_groups.append({
                        "label": label,
                        "spec_columns": spec_columns,
                        "rows": rendered_items,
                    })
                products_by_category[cat][page] = rendered_groups

    # 통계치 (히어로 지표) — 전부 실제 데이터에서 계산, 임의 수치 없음
    founded_year = ""
    iso_label = ""
    if history:
        founded_year = history[0].get("연월", "")[:4]
        for h in reversed(history):
            if "ISO 9001" in h.get("내용", ""):
                iso_label = "ISO 9001"
                break
    spec_count = 0
    if "제품상세" in wb.sheetnames:
        spec_count = sum(1 for r in wb["제품상세"].iter_rows(min_row=2, values_only=True) if any(r))
    stats = {
        "founded_year": founded_year,
        # Others(기타 브랜드 묶음)는 파트너사가 아니므로 제외
        "partner_count": sum(1 for a in business_areas if a.get("슬러그", "") != "Others"),
        "spec_count": spec_count,
        "iso_label": iso_label,
    }

    return {
        "company": company,
        "history": history,
        "org_by_dept": org_by_dept,
        "stats": stats,
        "business_areas": business_areas,
        "area_by_slug": area_by_slug,
        "brand_intro": brand_intro,
        "products_by_category": products_by_category,
    }


def build():
    safe_rmtree(DOCS_DIR)
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copytree(ROOT / "static", DOCS_DIR / "static", dirs_exist_ok=True)

    env = Environment(loader=FileSystemLoader(TEMPLATES_DIR))
    data = load_content()

    # 각 브랜드(사업영역)의 제품군 목록(슬러그 포함)을 미리 계산 -> 메뉴/브랜드페이지에서 공용으로 사용
    brand_slug_of = {a["슬러그"]: slugify(a["슬러그"]) for a in data["business_areas"]}
    family_slugs = {}  # cat -> {page_name: family_slug}
    family_summary = {}  # cat -> "취급 제품: A, B, C"
    for cat, pages in data["products_by_category"].items():
        family_slugs[cat] = {page: slugify(page) for page in pages}
        manual_summary = data["brand_intro"].get(cat, {}).get("제품요약", "")
        family_summary[cat] = manual_summary if manual_summary else ", ".join(pages.keys())

    common = dict(
        company=data["company"],
        business_areas=data["business_areas"],
        area_by_slug=data["area_by_slug"],
        products_by_category=data["products_by_category"],
        brand_intro=data["brand_intro"],
        brand_slug_of=brand_slug_of,
        family_slugs=family_slugs,
        family_summary=family_summary,
        stats=data["stats"],
        slugify=slugify,
        meta_keywords="케미프렌드,Chemifriend,Cabot,Carbon black,Synthomer,Arkema,Syensqo,화학원료 유통",
    )

    def render(template_name, out_path: Path, base_path: str, active: str, title: str, description: str, **extra):
        template = env.get_template(template_name)
        rel_path = out_path.relative_to(DOCS_DIR).as_posix()
        html = template.render(
            base_path=base_path,
            active=active,
            page_title=title,
            meta_description=description,
            canonical_url=f"{SITE_URL}/{rel_path}",
            **common,
            **extra,
        )
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(html, encoding="utf-8")
        sitemap_urls.append(rel_path)

    sitemap_urls = []
    SITE_URL = "https://chemifriend.com"

    # 1) 메인
    render("index.html", DOCS_DIR / "index.html", "", "main", "Main",
           f"{data['company'].get('회사명','')} 공식 홈페이지. {data['company'].get('슬로건','')}",
           history=data["history"], org_by_dept=data["org_by_dept"])

    # 2) Product 허브
    render("product_hub.html", DOCS_DIR / "product.html", "", "product", "Product",
           "케미프렌드 취급 브랜드 안내")

    # 3) 브랜드별 페이지 + 제품군별 개별 페이지
    for area in data["business_areas"]:
        slug = area["슬러그"]
        bslug = brand_slug_of[slug]
        pages = data["products_by_category"].get(slug, OrderedDict())
        intro = data["brand_intro"].get(slug, {})

        render(
            "brand.html", DOCS_DIR / "product" / bslug / "index.html", "../../", "product",
            area["회사명"],
            f"{area['회사명']} 취급 제품군 안내",
            area=area, intro=intro, pages=pages, fam_slugs=family_slugs.get(slug, {}),
        )

        for page_name, groups in pages.items():
            fslug = family_slugs[slug][page_name]
            render(
                "product_family.html", DOCS_DIR / "product" / bslug / fslug / "index.html", "../../../", "product",
                f"{page_name} — {area['회사명']}",
                f"{area['회사명']} {page_name} 제품 스펙",
                area=area, page_name=page_name, groups=groups, fslug=fslug,
            )
            for group in groups:
                for item in group["rows"]:
                    render(
                        "product_sku.html", DOCS_DIR / "product" / bslug / fslug / f"{item['sku_slug']}.html", "../../../", "product",
                        f"{item['품명']} — {area['회사명']} {page_name}",
                        f"{item['품명']} 스펙 — {area['회사명']} {page_name}",
                        area=area, page_name=page_name, fslug=fslug, group=group, item=item,
                    )

    # 4) Contact
    render("contact.html", DOCS_DIR / "contact.html", "", "contact", "Contact",
           "케미프렌드 오시는 길 및 문의처 안내")

    (DOCS_DIR / "CNAME").write_text("chemifriend.com\n", encoding="utf-8")

    # SEO: sitemap.xml + robots.txt 자동 생성
    sitemap_entries = "\n".join(
        f'  <url><loc>{SITE_URL}/{u}</loc></url>' for u in sitemap_urls
    )
    sitemap_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
        f"{sitemap_entries}\n"
        "</urlset>\n"
    )
    (DOCS_DIR / "sitemap.xml").write_text(sitemap_xml, encoding="utf-8")
    (DOCS_DIR / "robots.txt").write_text(
        f"User-agent: *\nAllow: /\nSitemap: {SITE_URL}/sitemap.xml\n", encoding="utf-8"
    )

    print(f"빌드 완료 → {DOCS_DIR} (페이지 {len(sitemap_urls)}개, sitemap.xml 포함)")


if __name__ == "__main__":
    build()
